#!/usr/bin/env python3
"""doc_kit — versioned documents (contracts, letters, schedules, decks) for the Cascade Timber corpus.

The attachment-builder agent writes document *content* as JSON in library/<DOC_ID>.json, with a
version history (draft → executed → scanned copy …), and an attach plan that maps corpus emails
to the version they should carry. This tool renders, validates and wires them in:

  list                               documents in the library and their versions
  render  <DOC_ID> [--version V]     render one version to Logs/attachments/preview/ (+ page-1 PNG)
  find    <regex>                    emails whose attachment filenames/subjects match (candidates)
  context <DOCID>                    one email + its current attachments' text + manifest row
  validate <plan.json>               checks (dates, cutoffs, parcels, figures, seed lock, mismatches)
  apply    <plan.json> [--dry-run]   rebuild those emails with the rendered attachments
  rollback <PLAN_ID>                 restore the emails from backup
  publish [--type contract]          export every version of library docs of that type (plus
                                     contract PDFs already attached to seed emails) to
                                     data/contracts/ with INDEX.csv — the standalone collection

Formats: pdf (reportlab), docx (python-docx), xlsx (openpyxl), csv, txt. Any PDF version can be
`"scan": true` — rasterized, slightly skewed, grey, no text layer — like a signed copy that was
printed, signed and scanned back in. Signatures: "wet" (hand-drawn look), "docusign", "typed" (/s/).
The data-room builder reuses render_version() and the library.
"""
import argparse
import copy
import csv
import io
import json
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "thread_kit"))
import thread_kit as tk  # noqa: E402

LIB = os.path.join(HERE, "library")
PLANS = os.path.join(HERE, "plans")
WORK = os.path.join(tk.ROOT, "Logs", "attachments")
LOGOS = os.path.join(HERE, "..", "thread_kit", "logos")
FONTS = "/System/Library/Fonts/Supplemental"
LETTERHEAD_LOGO = {
    "cascadetimber.com": "Cascade_Timber_logo.png",
    "llassociates.com": "LL_Associates_logo.png",
    "mosslane-cpa.com": "Moss_Lane_logo.png",
    "whitakerforensic.com": "Whitaker_Forensic_logo.png",
}
MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv", "txt": "text/plain",
}


# ---------------------------------------------------------------- library

def load_doc(doc_id):
    p = os.path.join(LIB, f"{doc_id}.json")
    if not os.path.exists(p):
        sys.exit(f"no library document {doc_id} ({p})")
    return json.load(open(p))


def all_docs():
    return {os.path.splitext(f)[0]: json.load(open(os.path.join(LIB, f)))
            for f in sorted(os.listdir(LIB)) if f.endswith(".json")} if os.path.isdir(LIB) else {}


def subst(obj, vars_):
    """Replace {{name}} placeholders everywhere in a JSON-like structure."""
    if isinstance(obj, str):
        def rep(m):
            k = m.group(1).strip()
            if k not in vars_:
                raise KeyError(f"undefined variable {{{{{k}}}}}")
            return str(vars_[k])
        return re.sub(r"\{\{\s*([\w.]+)\s*\}\}", rep, obj)
    if isinstance(obj, list):
        return [subst(x, vars_) for x in obj]
    if isinstance(obj, dict):
        return {k: subst(v, vars_) for k, v in obj.items()}
    return obj


def resolve(doc, version):
    """Merge base content + version: vars cascade (base vars < earlier versions < this one)."""
    vers = doc["versions"]
    names = [v["version"] for v in vers]
    if version is None:
        version = names[-1]
    if version not in names:
        raise KeyError(f"{doc['doc_id']} has no version {version} (has {names})")
    vars_ = dict(doc.get("vars", {}))
    for v in vers:
        vars_.update(v.get("vars", {}))
        if v["version"] == version:
            break
    ver = next(v for v in vers if v["version"] == version)
    content = copy.deepcopy(doc["content"])
    for k, v in ver.get("content_overrides", {}).items():
        content[k] = v
    content = subst(content, vars_)
    meta = {k: doc.get(k) for k in ("doc_id", "title", "type", "org", "parcels", "arc", "doc_control")}
    meta.update({k: ver.get(k) for k in ("version", "date", "status", "format", "scan", "filename",
                                         "signatures", "watermark", "changes", "sig_style")})
    meta["signatures"] = subst(ver.get("signatures", []), vars_)
    meta["filename"] = subst(ver.get("filename") or f"{doc['doc_id']}_{version}.{ver.get('format', 'pdf')}", vars_)
    meta["title"] = subst(doc.get("title", ""), vars_)
    meta["doc_control"] = subst(doc.get("doc_control") or "", vars_)
    meta["vars"] = vars_
    return meta, content


def doc_text(meta, content):
    """All human-readable text of a resolved version (for validation scans)."""
    out = [meta.get("title", "")]

    def walk(x):
        if isinstance(x, str):
            out.append(x)
        elif isinstance(x, list):
            for y in x:
                walk(y)
        elif isinstance(x, dict):
            for y in x.values():
                walk(y)
    walk(content)
    for s in meta.get("signatures") or []:
        out.append(" ".join(str(v) for v in s.values()))
    return "\n".join(out)


# ---------------------------------------------------------------- PDF rendering (reportlab)

def _fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    reg = pdfmetrics.getRegisteredFontNames()
    for name, fn in (("Sig", "Bradley Hand Bold.ttf"), ("Sig2", "Brush Script.ttf"),
                     ("Serif", "Georgia.ttf"), ("Serif-Bold", "Georgia Bold.ttf"),
                     ("Serif-Italic", "Georgia Italic.ttf")):
        if name not in reg:
            try:
                pdfmetrics.registerFont(TTFont(name, os.path.join(FONTS, fn)))
            except Exception:
                pass
    return "Serif" if "Serif" in pdfmetrics.getRegisteredFontNames() else "Times-Roman"


def render_pdf(meta, content):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)
    from xml.sax.saxutils import escape as X

    base = _fonts()
    bold = "Serif-Bold" if base == "Serif" else "Times-Bold"
    ital = "Serif-Italic" if base == "Serif" else "Times-Italic"
    body = ParagraphStyle("b", fontName=base, fontSize=10.2, leading=14, alignment=TA_JUSTIFY, spaceAfter=6)
    head = ParagraphStyle("h", parent=body, fontName=bold, fontSize=10.4, spaceBefore=6, spaceAfter=3, alignment=0)
    title = ParagraphStyle("t", parent=body, fontName=bold, fontSize=14, alignment=TA_CENTER, spaceAfter=4, leading=18)
    sub = ParagraphStyle("s", parent=body, fontName=ital, fontSize=9.5, alignment=TA_CENTER, spaceAfter=12)
    small = ParagraphStyle("sm", parent=body, fontSize=8.5, leading=11)
    buf = io.BytesIO()
    doc_ctl = meta.get("doc_control") or ""
    status = (meta.get("status") or "").lower()
    wm = meta.get("watermark") if meta.get("watermark") is not None else ("DRAFT" if status == "draft" else None)

    def page(c, d):
        c.saveState()
        c.setFont(base, 7.5)
        c.setFillColor(colors.HexColor("#555555"))
        left = f"{doc_ctl}  ·  {meta.get('version', '')}".strip(" ·")
        c.drawString(0.9 * inch, 0.55 * inch, left)
        c.drawRightString(LETTER[0] - 0.9 * inch, 0.55 * inch, f"Page {d.page}")
        if content.get("footer"):
            c.drawCentredString(LETTER[0] / 2, 0.38 * inch, content["footer"][:140])
        if wm:
            c.setFont(bold, 88)
            c.setFillColor(colors.Color(0.75, 0.75, 0.75, alpha=0.28))
            c.translate(LETTER[0] / 2, LETTER[1] / 2)
            c.rotate(38)
            c.drawCentredString(0, 0, wm)
        c.restoreState()

    story = []
    org = (meta.get("org") or "").lower()
    logo = LETTERHEAD_LOGO.get(org)
    if logo and os.path.exists(os.path.join(LOGOS, logo)):
        from PIL import Image as PI
        im = PI.open(os.path.join(LOGOS, logo))
        w = 2.1 * inch if im.width / im.height > 2.5 else 1.1 * inch
        story.append(Image(os.path.join(LOGOS, logo), width=w, height=w * im.height / im.width, hAlign="LEFT"))
        story.append(Spacer(1, 6))
    elif content.get("letterhead"):
        lh = content["letterhead"]
        story.append(Paragraph(f'<font name="{bold}" size="12">{X(lh.get("name", ""))}</font><br/>'
                               f'<font size="8.5">{X(lh.get("lines", ""))}</font>', ParagraphStyle("lh", parent=body, alignment=0)))
        story.append(Spacer(1, 10))
    story.append(Paragraph(X(meta.get("title", "")), title))
    if content.get("subtitle"):
        story.append(Paragraph(X(content["subtitle"]), sub))
    for p in content.get("preamble", []):
        story.append(Paragraph(X(p), body))
    if content.get("recitals"):
        story.append(Paragraph("RECITALS", head))
        for i, r in enumerate(content["recitals"]):
            story.append(Paragraph(f"{chr(65 + i)}. {X(r)}", body))
        if content.get("recitals_close"):
            story.append(Paragraph(X(content["recitals_close"]), body))
    for n, s in enumerate(content.get("sections", []), 1):
        num = s.get("number", f"{n}.")
        if s.get("heading"):
            story.append(Paragraph(f"{X(str(num))} {X(s['heading'])}", head))
        for para in ([s["text"]] if isinstance(s.get("text"), str) else s.get("text", [])):
            story.append(Paragraph(X(para), body))
        for item in s.get("items", []):
            story.append(Paragraph(X(item), ParagraphStyle("i", parent=body, leftIndent=18)))
        if s.get("table"):
            story.append(_pdf_table(s["table"], base, bold, Table, TableStyle, colors, small))
            story.append(Spacer(1, 6))
    if content.get("closing"):
        story.append(Paragraph(X(content["closing"]), body))
    sigs = meta.get("signatures") or []
    if sigs:
        story.append(Spacer(1, 10))
        story.append(Paragraph(X(content.get("signature_intro", "IN WITNESS WHEREOF, the parties have executed this Agreement as of the Effective Date.")), body))
        story.append(Spacer(1, 8))
        blocks = [_sig_block(s, meta.get("sig_style") or "wet", base, bold, Paragraph, ParagraphStyle, body, Table, TableStyle, colors) for s in sigs]
        rows = [blocks[i:i + 2] + ([""] if len(blocks[i:i + 2]) == 1 else []) for i in range(0, len(blocks), 2)]
        t = Table(rows, colWidths=[3.2 * inch, 3.2 * inch])
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        story.append(KeepTogether(t))
    for ex in content.get("exhibits", []):
        story.append(Spacer(1, 16))
        story.append(Paragraph(X(ex.get("title", "")), head))
        for para in ([ex["text"]] if isinstance(ex.get("text"), str) else ex.get("text", [])):
            story.append(Paragraph(X(para), body))
        if ex.get("table"):
            story.append(_pdf_table(ex["table"], base, bold, Table, TableStyle, colors, small))
    SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                      topMargin=0.8 * inch, bottomMargin=0.9 * inch,
                      title=meta.get("title", ""), author=content.get("pdf_author", ""),
                      subject=meta.get("doc_id", "")).build(story, onFirstPage=page, onLaterPages=page)
    return buf.getvalue()


def _pdf_table(rows, base, bold, Table, TableStyle, colors, small):
    from reportlab.platypus import Paragraph
    from xml.sax.saxutils import escape as X
    data = [[Paragraph(X(str(c)), small) for c in r] for r in rows]
    t = Table(data, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _sig_block(s, style, base, bold, Paragraph, ParagraphStyle, body, Table, TableStyle, colors):
    from xml.sax.saxutils import escape as X
    from reportlab.pdfbase import pdfmetrics
    lines = []
    party = ParagraphStyle("p", parent=body, fontName=bold, fontSize=9.5, leading=12, alignment=0)
    txt = ParagraphStyle("x", parent=body, fontSize=9.5, leading=12, alignment=0)
    lines.append(Paragraph(X(s.get("party", "")), party))
    signed = s.get("signed", False)
    sig_font = "Sig" if "Sig" in pdfmetrics.getRegisteredFontNames() else "Times-Italic"
    if signed and style == "docusign":
        lines.append(Paragraph(f'<font size="6.5" color="#1F4E9A">DocuSigned by:</font><br/>'
                               f'<font name="{sig_font}" size="16" color="#1F2A6B">{X(s.get("name", ""))}</font><br/>'
                               f'<font size="6" color="#1F4E9A">{X(s.get("envelope", "4F2A9C1E7B3D48A6"))}</font>', txt))
    elif signed and style == "typed":
        lines.append(Paragraph(f"/s/ {X(s.get('name', ''))}", txt))
    elif signed:
        lines.append(Paragraph(f'<font name="{sig_font}" size="17" color="#1B2A5A">{X(s.get("name", ""))}</font>', txt))
    else:
        lines.append(Paragraph("<br/>", txt))
    lines.append(Paragraph("By: ______________________________", txt))
    lines.append(Paragraph(f"Name: {X(s.get('name', ''))}", txt))
    if s.get("title"):
        lines.append(Paragraph(f"Title: {X(s['title'])}", txt))
    lines.append(Paragraph(f"Date: {X(s.get('date', '') if signed else '________________')}", txt))
    return lines


def scanify(pdf_bytes, seed):
    """Print-and-scan look: rasterize, greyscale, slight skew and noise; image-only PDF."""
    from PIL import Image, ImageFilter
    rnd = random.Random(seed)
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "in.pdf")
        open(src, "wb").write(pdf_bytes)
        subprocess.run(["pdftoppm", "-r", "110", "-gray", "-png", src, os.path.join(td, "p")], check=True)
        pages = []
        for fn in sorted(f for f in os.listdir(td) if f.startswith("p")):
            im = Image.open(os.path.join(td, fn)).convert("L")
            im = im.rotate(rnd.uniform(-0.9, 0.9), resample=Image.BICUBIC, expand=False, fillcolor=255)
            px = im.load()
            w, h = im.size
            for _ in range(w * h // 180):
                x, y = rnd.randrange(w), rnd.randrange(h)
                px[x, y] = max(0, px[x, y] - rnd.randint(40, 120))
            im = im.filter(ImageFilter.GaussianBlur(0.35)).point(lambda p: 245 if p > 225 else int(p * 0.97))
            pages.append(im)
        out = io.BytesIO()
        pages[0].save(out, "PDF", resolution=110, save_all=True, append_images=pages[1:])
        return out.getvalue()


# ---------------------------------------------------------------- DOCX / XLSX / CSV

def render_docx(meta, content):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(11)
    sec = d.sections[0]
    status = (meta.get("status") or "").lower()
    wm = meta.get("watermark") if meta.get("watermark") is not None else ("DRAFT" if status == "draft" else None)
    hdr = sec.header.paragraphs[0]
    hdr.text = " — ".join(x for x in [wm, meta.get("doc_control"), meta.get("version")] if x)
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if content.get("footer"):
        sec.footer.paragraphs[0].text = content["footer"]
    logo = LETTERHEAD_LOGO.get((meta.get("org") or "").lower())
    if logo and os.path.exists(os.path.join(LOGOS, logo)):
        d.add_picture(os.path.join(LOGOS, logo), width=Inches(2.0))
    elif content.get("letterhead"):
        p = d.add_paragraph()
        p.add_run(content["letterhead"].get("name", "")).bold = True
        p.add_run("\n" + content["letterhead"].get("lines", "")).font.size = Pt(8.5)
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(meta.get("title", ""))
    r.bold, r.font.size = True, Pt(14)
    if content.get("subtitle"):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(content["subtitle"]).italic = True
    for x in content.get("preamble", []):
        d.add_paragraph(x)
    if content.get("recitals"):
        d.add_paragraph().add_run("RECITALS").bold = True
        for i, x in enumerate(content["recitals"]):
            d.add_paragraph(f"{chr(65 + i)}. {x}")
        if content.get("recitals_close"):
            d.add_paragraph(content["recitals_close"])
    for n, s in enumerate(content.get("sections", []), 1):
        if s.get("heading"):
            d.add_paragraph().add_run(f"{s.get('number', f'{n}.')} {s['heading']}").bold = True
        for para in ([s["text"]] if isinstance(s.get("text"), str) else s.get("text", [])):
            d.add_paragraph(para)
        for item in s.get("items", []):
            d.add_paragraph(item, style="List Bullet")
        if s.get("table"):
            _docx_table(d, s["table"])
    if content.get("closing"):
        d.add_paragraph(content["closing"])
    sigs = meta.get("signatures") or []
    if sigs:
        d.add_paragraph(content.get("signature_intro", "IN WITNESS WHEREOF, the parties have executed this Agreement as of the Effective Date."))
        for s in sigs:
            p = d.add_paragraph()
            p.add_run(s.get("party", "")).bold = True
            if s.get("signed"):
                sr = p.add_run("\n" + s.get("name", ""))
                sr.font.name, sr.font.size = "Bradley Hand", Pt(16)
                sr.font.color.rgb = RGBColor(0x1B, 0x2A, 0x5A)
            p.add_run(f"\nBy: ______________________\nName: {s.get('name', '')}\nTitle: {s.get('title', '')}\n"
                      f"Date: {s.get('date', '') if s.get('signed') else '____________'}")
    for ex in content.get("exhibits", []):
        d.add_page_break()
        d.add_paragraph().add_run(ex.get("title", "")).bold = True
        for para in ([ex["text"]] if isinstance(ex.get("text"), str) else ex.get("text", [])):
            d.add_paragraph(para)
        if ex.get("table"):
            _docx_table(d, ex["table"])
    cp = d.core_properties
    cp.title, cp.author = meta.get("title", ""), content.get("docx_author", "")
    cp.last_modified_by = content.get("docx_author", "")
    if meta.get("date"):
        dt = datetime.fromisoformat(meta["date"])
        cp.created = cp.modified = dt
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_table(d, rows):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            t.cell(i, j).text = str(c)
            if i == 0:
                for run in t.cell(i, j).paragraphs[0].runs:
                    run.bold = True


def render_xlsx(meta, content):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    sheets = content.get("sheets") or [{"name": "Sheet1", "rows": content.get("table", [])}]
    for sh in sheets:
        ws = wb.create_sheet(sh.get("name", "Sheet")[:31])
        for r in sh.get("rows", []):
            ws.append(r)
        for c in ws[1] if ws.max_row else []:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8E8E8")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col) + 2))
    wb.properties.title = meta.get("title", "")
    wb.properties.creator = content.get("xlsx_author", "")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_csv(meta, content):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    for r in content.get("table", []):
        w.writerow(r)
    return buf.getvalue().encode("utf-8")


def render_version(doc, version=None):
    """-> (filename, mime type, bytes, meta)."""
    meta, content = resolve(doc, version)
    fmt = (meta.get("format") or "pdf").lower()
    if fmt == "pdf":
        data = render_pdf(meta, content)
        if meta.get("scan"):
            data = scanify(data, seed=f"{meta['doc_id']}{meta['version']}")
    elif fmt == "docx":
        data = render_docx(meta, content)
    elif fmt == "xlsx":
        data = render_xlsx(meta, content)
    elif fmt == "csv":
        data = render_csv(meta, content)
    elif fmt == "txt":
        data = content.get("text", "").encode("utf-8")
    else:
        raise ValueError(f"unsupported format {fmt}")
    return meta["filename"], MIME[fmt], data, meta


# ---------------------------------------------------------------- validation

DATE_PATTERNS = [
    (r"\b(January|February|March|April|May|June|July|August|September|October|November|December) (\d{1,2}), (20\d\d)\b", "%B %d %Y"),
    (r"\b(\d{1,2})/(\d{1,2})/(20\d\d|\d\d)\b", None),
    (r"\b(20\d\d)-(\d\d)-(\d\d)\b", "%Y-%m-%d"),
]


def dates_in(text):
    out = []
    for pat, fmt in DATE_PATTERNS:
        for m in re.finditer(pat, text):
            try:
                if fmt == "%B %d %Y":
                    out.append((m.group(0), datetime.strptime(f"{m.group(1)} {m.group(2)} {m.group(3)}", fmt).date()))
                elif fmt is None:
                    y = int(m.group(3))
                    y = y + 2000 if y < 100 else y
                    out.append((m.group(0), date(y, int(m.group(1)), int(m.group(2)))))
                else:
                    out.append((m.group(0), datetime.strptime(m.group(0), fmt).date()))
            except ValueError:
                pass
    return out


def validate_version(doc, version, issues, where):
    err = lambda m: issues.append(("ERROR", where, m))
    warn = lambda m: issues.append(("WARN", where, m))
    try:
        meta, content = resolve(doc, version)
    except KeyError as e:
        err(str(e))
        return None
    vdate = datetime.fromisoformat(meta["date"]).date() if meta.get("date") else None
    if not vdate:
        err("version has no date")
        return meta
    text = doc_text(meta, content)
    dt = datetime.combine(vdate, datetime.min.time()).replace(tzinfo=tk.TZ)
    for h in tk.in_period(tk.RULES["terms"], text, dt):
        err(f"knowledge cutoff in document text: /{h['pattern']}/ not before {h['not_before']} ({h['anchor']})")
    for h in tk.RULES["never"]:
        if re.search(h["pattern"], text, re.I):
            err(f"banned by design: {h['reason']}")
    kw = any(p.startswith("KW") for p in (meta.get("parcels") or [])) or re.search(r"\bKW-0[12]\b|Clearwater", text)
    for s, d in dates_in(text):
        if kw and d <= date(2022, 3, 16):
            err(f"KW/Clearwater document shows date {s} on or before 3/16/22 (Arc J timeline guard)")
    if kw and vdate <= date(2022, 3, 16):
        err("KW/Clearwater document version dated on or before 3/16/22")
    for s in meta.get("signatures") or []:
        if s.get("signed"):
            if not s.get("date"):
                err(f"signed block for {s.get('name')} has no date")
            else:
                sd = next((d for _, d in dates_in(s["date"])), None)
                if sd and sd > vdate:
                    err(f"{s.get('name')} signs {sd} — after the version date {vdate}")
    if (meta.get("status") or "").lower() == "executed" and not all(s.get("signed") for s in meta.get("signatures") or [{}]):
        warn("status executed but not every signature block is signed (partially executed?)")
    known = set(tk.RULES["known_figures"]) | set(doc.get("figures_ok", []))
    for fig in re.findall(r"\$\s?\d[\d,]*(?:\.\d+)?(?:\s?[KMB]\b)?", text):
        fig = fig.rstrip(",. ")
        if fig.replace(" ", "") not in known and fig not in ("$0", "$0."):
            warn(f"figure {fig} not in bible ledger or doc figures_ok")
    for pid in set(re.findall(r"\b[A-Z]{2}-\d{2}\b", text)):
        if pid not in tk.RULES["known_parcels"] and not pid.startswith("RM"):
            warn(f"parcel id {pid} not in the bible")
    return meta


def validate_plan(plan, docs, library):
    issues = []
    err = lambda k, m: issues.append(("ERROR", k, m))
    warn = lambda k, m: issues.append(("WARN", k, m))
    if not plan.get("plan_id"):
        err("-", "plan needs plan_id")
    if not plan.get("bible_refs"):
        warn("-", "plan should cite bible_refs")
    for doc_id in {a["doc"].split("@")[0] for a in plan.get("attach", [])}:
        if doc_id not in library:
            err("-", f"library document {doc_id} missing")
            continue
        doc = library[doc_id]
        dates = [v.get("date") for v in doc["versions"]]
        if dates != sorted(dates):
            err(doc_id, "versions are not in date order")
        for v in doc["versions"]:
            validate_version(doc, v["version"], issues, f"{doc_id}@{v['version']}")
    seen = set()
    for a in plan.get("attach", []):
        k = a["email"]
        if k in seen:
            err(k, "email listed twice (put multiple attachments in one entry via 'docs')")
        seen.add(k)
        d = docs.get(k)
        if not d:
            err(k, "email not in corpus")
            continue
        if tk.num(k) <= tk.SEED_MAX:
            err(k, "seed email — locked; attach via a new later email instead")
        doc_id, _, ver = a["doc"].partition("@")
        if doc_id not in library:
            continue
        try:
            meta, _ = resolve(library[doc_id], ver or None)
        except KeyError as e:
            err(k, str(e))
            continue
        edate = d.date.astimezone(tk.TZ).date()
        if datetime.fromisoformat(meta["date"]).date() > edate:
            err(k, f"{a['doc']} is dated {meta['date']} — after this email ({edate})")
        for s in meta.get("signatures") or []:
            sd = next((x for _, x in dates_in(s.get("date", ""))), None) if s.get("signed") else None
            if sd and sd > edate:
                err(k, f"attachment shows a signature dated {sd}, after the email ({edate})")
        names = [x[0] for x in d.attachments]
        if a.get("replace") and a["replace"] not in names:
            err(k, f"replace target {a['replace']!r} not on email (has {names})")
        if not a.get("replace") and not a.get("add"):
            err(k, "entry needs 'replace': <filename> or 'add': true")
        # flag pre-existing email defects the reader will trip over
        mentioned = set(re.findall(r"\bKW-0[12]\b", d.new_text)) | set(re.findall(r"\bKW-0[12]\b", d.subject))
        own = set(meta.get("parcels") or [])
        if mentioned and own and not (mentioned & own):
            warn(k, f"email mentions {sorted(mentioned)} but the attachment is for {sorted(own)}")
        elif mentioned and own and not (set(re.findall(r'\bKW-0[12]\b', d.subject)) <= own) and set(re.findall(r'\bKW-0[12]\b', d.subject)):
            warn(k, f"pre-existing defect: subject names {re.findall(r'KW-0[12]', d.subject)} but body/attachment are {sorted(own)}")
        if d.from_addr in {a2 for _, a2 in d.to} and not d.cc:
            warn(k, "pre-existing defect: email sent to self")
        if a.get("note"):
            issues.append(("INFO", k, a["note"]))
    return issues


# ---------------------------------------------------------------- apply / rollback

def apply(plan, docs, library, dry):
    issues = validate_plan(plan, docs, library)
    tk.report_issues(issues)
    if any(i[0] == "ERROR" for i in issues):
        sys.exit("apply aborted: fix ERRORs first")
    pid = plan["plan_id"]
    bdir = os.path.join(WORK, "backup", pid)
    manifest, fields = tk.load_manifest()
    rendered = {}
    rec = {"plan_id": pid, "emails": [], "applied_at": datetime.now().isoformat(timespec="seconds")}
    for a in plan["attach"]:
        d = docs[a["email"]]
        key = a["doc"]
        if key not in rendered:
            doc_id, _, ver = key.partition("@")
            rendered[key] = render_version(library[doc_id], ver or None)
        fn, mime, data, meta = rendered[key]
        fn = a.get("filename", fn)
        atts = [x for x in d.attachments if x[0] != a.get("replace")] if a.get("replace") else list(d.attachments)
        idx = next((i for i, x in enumerate(d.attachments) if x[0] == a.get("replace")), len(atts))
        atts.insert(min(idx, len(atts)), (fn, mime, data))
        print(f"{a['email']}: {a.get('replace') or '(add)'} -> {fn} ({len(data) // 1024} KB, {meta['status']})")
        if dry:
            continue
        import email as E
        from email import policy
        with open(d.path, "rb") as f:
            m = E.message_from_binary_file(f, policy=policy.default)
        hdrs = [(k, v) for k, v in m.items() if not k.lower().startswith("content-") and k.lower() != "mime-version"]
        os.makedirs(bdir, exist_ok=True)
        shutil.copy2(d.path, os.path.join(bdir, os.path.basename(d.path)))
        with open(d.path, "wb") as f:
            f.write(tk.build_mime(hdrs, d.body, atts))
        rec["emails"].append({"email": a["email"], "path": os.path.relpath(d.path, tk.ROOT), "doc": key, "filename": fn})
        if a["email"] in manifest:
            manifest[a["email"]]["expected_attachment_count"] = str(len(atts))
    if dry:
        return
    with open(tk.MANIFEST, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(fields), lineterminator="\n")
        w.writeheader()
        w.writerows(manifest.values())
    os.makedirs(WORK, exist_ok=True)
    json.dump(rec, open(os.path.join(WORK, f"{pid}.apply.json"), "w"), indent=2)
    write_report(plan, rec, issues)
    print(f"\napplied {len(rec['emails'])} attachments; report Logs/attachments/{pid}_REPORT.md")


def write_report(plan, rec, issues):
    L = [f"# {plan['plan_id']} — attachment build report", "", plan.get("rationale", ""), "",
         f"Bible refs: {', '.join(plan.get('bible_refs', []))}", "",
         "| Email | Document@version | Filename |", "|---|---|---|"]
    L += [f"| {e['email']} | {e['doc']} | {e['filename']} |" for e in rec["emails"]]
    L += ["", "## Validation", ""] + ([f"- {l} [{k}] {m}" for l, k, m in issues if l != "ERROR"] or ["- clean"])
    open(os.path.join(WORK, f"{plan['plan_id']}_REPORT.md"), "w").write("\n".join(L) + "\n")


def rollback(pid):
    rec = json.load(open(os.path.join(WORK, f"{pid}.apply.json")))
    bdir = os.path.join(WORK, "backup", pid)
    for e in rec["emails"]:
        shutil.copy2(os.path.join(bdir, os.path.basename(e["path"])), os.path.join(tk.ROOT, e["path"]))
    print(f"restored {len(rec['emails'])} emails (manifest attachment counts unchanged by replace; re-check if you used 'add')")
    os.remove(os.path.join(WORK, f"{pid}.apply.json"))


# ---------------------------------------------------------------- publish (standalone collection)

SEED_CONTRACT_RE = re.compile(r"agreement|engagement|contract|option|amendment|letter", re.I)


def publish(library, docs, doc_type):
    out_root = os.path.join(tk.ROOT, "data", "contracts")
    carried, first_copy = {}, {}
    for rec_p in sorted(os.listdir(WORK)) if os.path.isdir(WORK) else []:
        if rec_p.endswith(".apply.json"):
            for e in json.load(open(os.path.join(WORK, rec_p)))["emails"]:
                carried.setdefault(e["doc"], []).append(e["email"])
                first_copy.setdefault(e["doc"], (e["email"], e["filename"]))
    rows = []
    for doc_id, doc in library.items():
        if doc_id.startswith("_") or (doc_type and doc.get("type") != doc_type):
            continue
        for v in doc["versions"]:
            key = f"{doc_id}@{v['version']}"
            fn, mime, data, meta = render_version(doc, v["version"])
            if key in first_copy and first_copy[key][0] in docs:
                # same bytes as the attached copy, so hash-dedup matches the email family
                em, efn = first_copy[key]
                hit = [x for x in docs[em].attachments if x[0] == efn]
                if hit:
                    fn, data = efn, hit[0][2]
            d = os.path.join(out_root, doc_id)
            os.makedirs(d, exist_ok=True)
            open(os.path.join(d, fn), "wb").write(data)
            rows.append({"doc_id": doc_id, "version": v["version"], "date": meta["date"], "status": meta.get("status", ""),
                         "title": meta["title"], "parcels": ";".join(meta.get("parcels") or []), "arc": meta.get("arc") or "",
                         "path": os.path.relpath(os.path.join(d, fn), out_root), "source": "doc_kit library",
                         "carried_by_emails": ";".join(sorted(set(carried.get(key, []))))})
    # contracts already attached to corpus emails outside the library (seed PDFs etc.)
    for dd in sorted(docs.values(), key=lambda x: tk.num(x.docid)):
        for fn, ct, data in dd.attachments:
            if ct == "application/pdf" and SEED_CONTRACT_RE.search(fn or "") and tk.num(dd.docid) <= tk.SEED_MAX:
                d = os.path.join(out_root, "seed")
                os.makedirs(d, exist_ok=True)
                open(os.path.join(d, fn), "wb").write(data)
                rows.append({"doc_id": f"SEED-{dd.docid}", "version": "as-attached",
                             "date": dd.date.astimezone(tk.TZ).date().isoformat(), "status": "as-attached",
                             "title": os.path.splitext(fn)[0].replace("_", " "), "parcels": "", "arc": "",
                             "path": os.path.relpath(os.path.join(d, fn), out_root), "source": f"attachment on {dd.docid} (seed, locked)",
                             "carried_by_emails": dd.docid})
    rows.sort(key=lambda r: (r["date"], r["doc_id"]))
    with open(os.path.join(out_root, "INDEX.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["doc_id"], lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    print(f"published {len(rows)} files to data/contracts/ (INDEX.csv)")


# ---------------------------------------------------------------- CLI

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("list")
    r = sub.add_parser("render"); r.add_argument("doc_id"); r.add_argument("--version")
    f = sub.add_parser("find"); f.add_argument("pattern")
    c = sub.add_parser("context"); c.add_argument("docid")
    v = sub.add_parser("validate"); v.add_argument("plan")
    a = sub.add_parser("apply"); a.add_argument("plan"); a.add_argument("--dry-run", action="store_true")
    rb = sub.add_parser("rollback"); rb.add_argument("plan_id")
    pb = sub.add_parser("publish"); pb.add_argument("--type", default="contract", help="library doc type to export ('' = all)")
    args = ap.parse_args()
    library = all_docs()
    if args.cmd == "list":
        for k, d in library.items():
            print(f"{k}: {resolve(d, None)[0]['title']}  [{', '.join(v['version'] + ' ' + v['date'] + ' ' + v.get('status', '') for v in d['versions'])}]")
        return
    if args.cmd == "render":
        fn, mime, data, meta = render_version(library[args.doc_id] if args.doc_id in library else load_doc(args.doc_id), args.version)
        issues = []
        validate_version(library[args.doc_id], meta["version"], issues, f"{args.doc_id}@{meta['version']}")
        tk.report_issues(issues)
        out = os.path.join(WORK, "preview")
        os.makedirs(out, exist_ok=True)
        p = os.path.join(out, fn)
        open(p, "wb").write(data)
        if fn.endswith(".pdf"):
            subprocess.run(["pdftoppm", "-r", "60", "-png", "-f", "1", "-l", "1", p, p[:-4] + "_p1"], check=False)
        print(f"wrote {os.path.relpath(p, tk.ROOT)} ({len(data) // 1024} KB)")
        return
    if args.cmd == "rollback":
        return rollback(args.plan_id)
    docs = tk.load_corpus()
    if args.cmd == "publish":
        return publish(library, docs, args.type)
    if args.cmd == "find":
        rx = re.compile(args.pattern, re.I)
        for d in sorted(docs.values(), key=lambda d: d.date):
            names = [x[0] or "" for x in d.attachments]
            if rx.search(d.subject) or any(rx.search(n) for n in names):
                print(f"{d.docid} {d.date.astimezone(tk.TZ):%Y-%m-%d} {d.from_name} → {', '.join(n or a for n, a in d.to)} | {d.subject} | {names}"
                      + ("  [SEED]" if tk.num(d.docid) <= tk.SEED_MAX else ""))
        return
    if args.cmd == "context":
        d = docs[args.docid.upper()]
        man, _ = tk.load_manifest()
        r = man.get(d.docid, {})
        print(f"{d.docid} {d.date.astimezone(tk.TZ):%a %Y-%m-%d %H:%M} {'[SEED — locked]' if tk.num(d.docid) <= tk.SEED_MAX else ''}")
        print(f"From {d.from_name} <{d.from_addr}> To {d.to} Cc {d.cc}\nSubject: {d.subject}")
        print(f"Manifest: arc={r.get('arc')} event={r.get('event_id')} role={r.get('intended_evidentiary_role')}\n")
        print(d.body[:3000])
        for fn, ct, data in d.attachments:
            print(f"\n--- attachment {fn} ({ct}, {len(data)} bytes)")
            if ct.startswith("text/"):
                print(data.decode("utf-8", "ignore")[:1500])
        return
    plan = json.load(open(args.plan))
    if args.cmd == "validate":
        issues = validate_plan(plan, docs, library)
        tk.report_issues(issues)
        sys.exit(1 if any(i[0] == "ERROR" for i in issues) else 0)
    if args.cmd == "apply":
        apply(plan, docs, library, args.dry_run)


if __name__ == "__main__":
    main()
