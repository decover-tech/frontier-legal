"""Reproducibly extract permitted Cascade source assets; no inference or network.

Run from the repository root: python3 -m tools.litigation_tasks.build_evidence
Requires pypdf, openpyxl, Pillow, pdftoppm and tesseract only at build time.
"""
import argparse
from collections import Counter
from datetime import timezone
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from hashlib import sha256 as _sha
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from .evidence import BUNDLE_PATH, VERSION, sha256

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
# Exact original paragraph allowlists, manually inspected against all four DOCXs.
# The summons is the instrument; the other protocols mix rules with answer keys.
PROTOCOL_RULE_PARAGRAPHS = {
    "Attorney_Client_Privilege.docx": set(range(1, 37)) | {37, 38, 40, 41, 42, 43, 44, 45, 46, 47, 48, 49, 54, 55, 59, 61, 66, 70, 71, 72, 73, 74, 75, 76, 77, 78, 79, 80, 81},
    "Responsiveness.docx": set(range(1, 27)) | {30, 31, 35, 36, 40, 41, 45, 46, 50, 53, 58, 59},
    "Work_Product.docx": set(range(1, 31)) | {49, 53, 55, 56},
}


def canonical(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def segments_to_text(segments):
    text, locators = "", []
    for label, content, extra in segments:
        content = content.replace("\r\n", "\n").replace("\r", "\n").strip()
        if not content:
            continue
        if text:
            text += "\n\n"
        start = len(text)
        text += content
        locators.append({"locator": label, "start": start, "end": len(text), **extra})
    return text, locators


def plain_segments(text, prefix="line"):
    return [(f"{prefix}:{i}", line, {"line": i}) for i, line in enumerate(text.splitlines(), 1) if line.strip()]


def docx_segments(raw):
    segments, warnings = [], []
    with ZipFile(BytesIO(raw)) as archive:
        names = [n for n in archive.namelist() if re.fullmatch(r"word/(document|footnotes|endnotes|comments|header\d+|footer\d+)\.xml", n)]
        names.sort(key=lambda n: (n != "word/document.xml", n))
        for name in names:
            xml = ET.fromstring(archive.read(name))
            for number, paragraph in enumerate(xml.findall(".//w:p", NS), 1):
                pieces = []
                for node in paragraph.iter():
                    if node.tag == "{" + NS["w"] + "}t":
                        pieces.append(node.text or "")
                    elif node.tag == "{" + NS["w"] + "}delText":
                        pieces.append("[DELETED: " + (node.text or "") + "]")
                if paragraph.findall(".//w:ins", NS) or paragraph.findall(".//w:del", NS):
                    warnings.append("Tracked revisions present; deletions marked, inserted text retained; inspect original for revision boundaries.")
                part = Path(name).stem
                segments.append((f"{part}:paragraph:{number}", "".join(pieces), {"paragraph": number, "part": part}))
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        if media:
            warnings.append(f"DOCX has {len(media)} embedded media asset(s); text extraction does not interpret images.")
    return segments, warnings


def ocr_image(path):
    if not shutil.which("tesseract"):
        raise RuntimeError("tesseract unavailable")
    result = subprocess.run(["tesseract", str(path), "stdout", "-l", "eng", "--psm", "6"], capture_output=True, timeout=120, check=True)
    return result.stdout.decode("utf-8").strip()


def extract(raw, mime):
    """Return extracted segments and honest text/OCR coverage, cached by payload.

    OCR output is a transcription candidate, never verified source truth. Empty
    OCR or unsupported content stays visible as incomplete extraction metadata.
    """
    segments, warnings, failures = [], [], []
    coverage = {"status": "complete_text", "pages": None, "native_text_pages": 0,
                "ocr_attempted_pages": 0, "ocr_text_pages": 0, "unextracted_pages": [],
                "visual_content_reviewed": False, "ocr_human_verified": False}
    try:
        if mime == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(raw))
            coverage["pages"] = len(reader.pages)
            with tempfile.TemporaryDirectory(prefix="cascade-evidence-") as directory:
                temp = Path(directory)
                (temp / "source.pdf").write_bytes(raw)
                for i, page in enumerate(reader.pages, 1):
                    content = page.extract_text() or ""
                    method = "native_text"
                    if content.strip():
                        coverage["native_text_pages"] += 1
                    else:
                        coverage["ocr_attempted_pages"] += 1
                        method = "ocr_unverified"
                        try:
                            if not shutil.which("pdftoppm"):
                                raise RuntimeError("pdftoppm unavailable")
                            subprocess.run(["pdftoppm", "-f", str(i), "-l", str(i), "-singlefile", "-r", "200", "-png", str(temp / "source.pdf"), str(temp / "page")], capture_output=True, timeout=120, check=True)
                            content = ocr_image(temp / "page.png")
                            if content:
                                coverage["ocr_text_pages"] += 1
                        except Exception as exc:
                            failures.append(f"page {i}: {type(exc).__name__}")
                        if not content.strip():
                            coverage["unextracted_pages"].append(i)
                    segments.append((f"page:{i}", content, {"page": i, "method": method}))
            warnings.append("PDF text does not establish signature authenticity, rendering fidelity, or complete image/annotation coverage.")
            if coverage["ocr_attempted_pages"]:
                coverage["status"] = "ocr_unverified"
                warnings.append("OCR may misread numbers, names, handwriting, or table structure; consult source page before relying on an exact disputed transcription.")
        elif mime == DOCX:
            segments, extra = docx_segments(raw)
            warnings.extend(extra)
            warnings.append("DOCX paragraph order includes table paragraphs; layout, signatures and pagination require original inspection.")
        elif mime == XLSX:
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(raw), data_only=False, read_only=True)
            for si, sheet in enumerate(workbook, 1):
                segments.append((f"sheet:{si}:title", sheet.title, {"sheet": si}))
                for row in sheet.iter_rows():
                    cells = [f"{c.coordinate}={c.value}" for c in row if c.value is not None]
                    if cells:
                        segments.append((f"sheet:{si}:row:{row[0].row}", " | ".join(cells), {"sheet": si, "row": row[0].row}))
            warnings.append("Cell values/formulas extracted; formulas are not recalculated and charts/images/formatting are not interpreted.")
        elif mime == PPTX:
            with ZipFile(BytesIO(raw)) as archive:
                names = sorted([n for n in archive.namelist() if re.fullmatch(r"ppt/(slides/slide|notesSlides/notesSlide)\d+\.xml", n)], key=lambda n: ("notes" in n, int(re.search(r"(\d+)\.xml$", n)[1])))
                for name in names:
                    xml = ET.fromstring(archive.read(name))
                    n = int(re.search(r"(\d+)\.xml$", name)[1])
                    kind = "notes" if "notes" in name else "slide"
                    content = "\n".join("".join(t.text or "" for t in p.findall(".//a:t", NS)) for p in xml.findall(".//a:p", NS))
                    segments.append((f"{kind}:{n}", content, {kind: n}))
            warnings.append("Slide and notes text extracted; figures, image text and spatial layout are not interpreted.")
        elif mime.startswith("image/"):
            from PIL import Image
            with tempfile.TemporaryDirectory(prefix="cascade-image-") as directory:
                path = Path(directory) / "image.png"
                image = Image.open(BytesIO(raw)); image.save(path)
                coverage["pixel_dimensions"] = list(image.size)
                content = ocr_image(path)
                segments = [("image:1", content, {"image": 1, "method": "ocr_unverified"})]
            coverage["status"] = "ocr_unverified" if content else "no_text_detected"
            warnings.append("Image OCR is not a determination that the image contains no evidence; inline logos/graphics retained as distinct payloads.")
        elif mime.startswith("text/"):
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1252")
                warnings.append("Decoded as Windows-1252 after UTF-8 failure.")
            segments = plain_segments(text)
        else:
            coverage["status"] = "unsupported"
            failures.append("Unsupported MIME type: " + mime)
    except Exception as exc:
        failures.append(type(exc).__name__)
    if failures or coverage["unextracted_pages"]:
        coverage["status"] = "incomplete"
    text, locators = segments_to_text(segments)
    if not text and coverage["status"] == "complete_text":
        coverage["status"] = "no_text_detected"
    return {"text": text, "locators": locators, "text_sha256": sha256(text.encode()),
            "extraction": {**coverage, "warnings": sorted(set(warnings)), "failures": failures}}


def build(root, output=None):
    root = Path(root).resolve()
    output = Path(output) if output else root / BUNDLE_PATH
    documents, provenance, source_files, cache, full_protocols = {}, {}, [], {}, {}
    counts = Counter()

    def extracted(raw, mime):
        cache_key = (sha256(raw), mime)
        if cache_key not in cache:
            cache[cache_key] = extract(raw, mime)
        return cache[cache_key]

    def add(key, raw, mime, kind, date, parent, path, part=None, filename=None, result=None):
        if key in documents:
            raise ValueError("Duplicate evidence ID " + key)
        result = result or extracted(raw, mime)
        documents[key] = {"document_id": key, "kind": kind, "mime_type": mime,
                          "source_sha256": sha256(raw), "source_date": date,
                          "parent_email_id": parent, **result}
        provenance[key] = {"source_path": path, "mime_part_index": part, "original_filename": filename,
                           "source_sha256": sha256(raw), "text_sha256": result["text_sha256"]}
        counts[kind] += 1

    for path in sorted(root.glob("data/emails/Custodians/**/*.eml")):
        raw = path.read_bytes(); relative = path.relative_to(root).as_posix()
        source_files.append({"path": relative, "sha256": sha256(raw)})
        key = re.match(r"EMAIL-\d+", path.name)[0]
        message = BytesParser(policy=policy.default).parsebytes(raw)
        if any(k.lower().startswith("x-decover-") for k in message.keys()):
            raise ValueError("Forbidden hidden label header in " + relative)
        date = parsedate_to_datetime(str(message["Date"]))
        if date.utcoffset() is None or len(message.get_all("Date", [])) != 1:
            raise ValueError("Ambiguous email date in " + relative)
        body = message.get_body(preferencelist=("plain",))
        if body is None:
            raise ValueError("No plain email body in " + relative)
        text = "\n".join(f"{h}: {message[h]}" for h in ("Date", "From", "To", "Cc", "Subject") if message[h]) + "\n\n" + body.get_content()
        result = extract(text.encode(), "text/plain")
        add(key, raw, "message/rfc822", "email", date.isoformat(), None, relative, result=result)
        attachments = []
        for index, part in enumerate(message.walk()):
            if part.is_multipart() or not (part.get_filename() or part.get_content_disposition() in ("attachment", "inline")):
                continue
            payload = part.get_payload(decode=True) or b""
            aid = f"ATT-{key}-{index:02d}-{sha256(payload)[:12]}"
            add(aid, payload, part.get_content_type(), "inline_image" if part.get_content_disposition() == "inline" else "attachment",
                None, key, relative, index, part.get_filename())
            attachments.append(aid)
        documents[key]["attachment_ids"] = attachments
    eligible = [*root.glob("definitions/*.docx"), *root.glob("data/contracts/**/*.pdf"), *root.glob("data/contracts/**/*.docx"), *root.glob("data/emails/Exhibits/**/*.pdf")]
    for path in sorted(eligible):
        raw = path.read_bytes(); relative = path.relative_to(root).as_posix()
        source_files.append({"path": relative, "sha256": sha256(raw)})
        # Include path hash in ID to preserve same-byte assets' distinct provenance.
        key = "DOC-" + sha256(relative.encode() + b"\0" + raw)[:20]
        kind = "protocol" if relative.startswith("definitions/") else "contract" if relative.startswith("data/contracts/") else "exhibit"
        mime = DOCX if path.suffix == ".docx" else "application/pdf"
        result = extracted(raw, mime)
        if path.name in PROTOCOL_RULE_PARAGRAPHS:
            full_protocols[key] = result
            original_segments, _ = docx_segments(raw)
            retained = [s for s in original_segments if s[2]["part"] == "document" and s[2]["paragraph"] in PROTOCOL_RULE_PARAGRAPHS[path.name]]
            omitted = [{"locator": s[0], "reason": "case-specific worked answer, answer-bearing example, branding, or adjoining example heading", "text_sha256": sha256(s[1].encode())} for s in original_segments if s[1].strip() and s not in retained]
            text, locators = segments_to_text(retained)
            result = {**result, "text": text, "locators": locators, "text_sha256": sha256(text.encode()),
                      "projection": {"type": "rule_only", "original_text_sha256": full_protocols[key]["text_sha256"], "omitted_paragraph_count": len(omitted), "omissions_are_extraction_failures": False}}
            provenance.setdefault(key, {})["projection_omissions"] = omitted
        omissions = provenance.get(key, {}).get("projection_omissions")
        add(key, raw, mime, kind, None, None, relative, filename=path.name, result=result)
        if omissions is not None:
            provenance[key]["projection_omissions"] = omissions
    coverage = {"version": VERSION, "document_count": len(documents), "counts_by_kind": dict(sorted(counts.items())),
                "counts_by_extraction_status": dict(sorted(Counter(d["extraction"]["status"] for d in documents.values()).items())),
                "unique_extracted_payloads": len(cache), "unique_pdf_pages": sum(r["extraction"]["pages"] or 0 for (h,m),r in cache.items() if m == "application/pdf"),
                "unique_pdf_native_text_pages": sum(r["extraction"]["native_text_pages"] for (h,m),r in cache.items() if m == "application/pdf"),
                "unique_pdf_ocr_attempted_pages": sum(r["extraction"]["ocr_attempted_pages"] for (h,m),r in cache.items() if m == "application/pdf"),
                "unique_pdf_ocr_text_pages": sum(r["extraction"]["ocr_text_pages"] for (h,m),r in cache.items() if m == "application/pdf"),
                "failure_document_ids": sorted(k for k,d in documents.items() if d["extraction"]["failures"]),
                "ocr_human_verified": False,
                "intentional_rule_projections": len(full_protocols),
                "limitations": ["No missing/empty extraction is evidence of absence.", "OCR transcriptions are unverified and must not decide disputed exact text without page review.", "Paragraph/page locators refer to the original container, not event time.", "Only email Date is machine assigned as source_date; attachment creation/execution dates remain null.", "DOCX media, spreadsheet charts and presentation images are not OCRed.", "Three protocols have rule-only public projections; full text and omitted-paragraph audit are evaluator-only manifest provenance. The summons instrument is complete.", "No filesystem indexes, authoring notes, hidden gold or analysis labels are included in public documents."]}
    output.mkdir(parents=True, exist_ok=True)
    data = b"".join((json.dumps(documents[k], sort_keys=True, ensure_ascii=False) + "\n").encode() for k in sorted(documents))
    (output / "documents.jsonl").write_bytes(data)
    (output / "coverage.json").write_bytes(canonical(coverage))
    import pypdf, openpyxl
    versions = {"pypdf": pypdf.__version__, "openpyxl": openpyxl.__version__}
    for binary, flag in (("tesseract", "--version"), ("pdftoppm", "-v")):
        if shutil.which(binary):
            p = subprocess.run([binary, flag], capture_output=True, text=True)
            versions[binary] = (p.stdout or p.stderr).splitlines()[0]
        else:
            versions[binary] = "unavailable"
    manifest = {"version": VERSION, "document_count": len(documents), "documents_sha256": sha256(data),
                "coverage_sha256": sha256((output / "coverage.json").read_bytes()), "extractor_versions": versions,
                "extractor_sha256": sha256(Path(__file__).read_bytes()), "source_files": source_files,
                "provenance": provenance, "full_protocol_extractions_evaluator_only": full_protocols,
                "ordering": "Lexicographic stable opaque document IDs; MIME index is depth-first walk index.",
                "public_document_contract": "documents.jsonl contains no source paths or original filenames as metadata. Faithful source text may naturally mention filenames. Manifest is trusted/evaluator-only."}
    (output / "manifest.json").write_bytes(canonical(manifest))
    return coverage


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    print(json.dumps(build(args.root, args.output), indent=2))
